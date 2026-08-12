import csv
import json
import os

from openpyxl import Workbook

from gas_sensor_daq.models.records import FIELDNAMES


BOOLEAN_FIELDS = {"humidity_on", "heating_on"}


def row_for_file(record):
    row = record.to_dict()
    normalized = {}
    for field in FIELDNAMES:
        value = row.get(field, "")
        if field in BOOLEAN_FIELDS:
            value = "true" if bool(value) else "false"
        normalized[field] = value
    return normalized


class CSVLogger:
    def __init__(self, filename, metadata=None):
        os.makedirs(os.path.dirname(filename) or "data", exist_ok=True)
        self.filename = filename
        self.metadata = dict(metadata or {})
        self.file = open(filename, "w", newline="")
        self.writer = csv.DictWriter(self.file, fieldnames=FIELDNAMES)
        self.writer.writeheader()
        self._write_metadata()

    def _write_metadata(self):
        metadata_filename = f"{os.path.splitext(self.filename)[0]}.metadata.json"
        with open(metadata_filename, "w", encoding="utf-8") as metadata_file:
            json.dump(self.metadata, metadata_file, ensure_ascii=False, indent=2)

    def update_metadata(self, updates):
        self.metadata.update(updates)
        self._write_metadata()

    def write(self, record):
        # FIELDNAMES controls column order and filters accidental extra keys.
        self.writer.writerow(row_for_file(record))
        self.file.flush()

    def close(self):
        self.file.close()


class ExcelLogger:
    def __init__(self, filename, metadata=None):
        os.makedirs(os.path.dirname(filename) or "data", exist_ok=True)
        self.filename = filename
        self.metadata = dict(metadata or {})
        self.workbook = Workbook()
        self.sheet = self.workbook.active
        self.sheet.title = "Data"
        self.sheet.append(FIELDNAMES)
        self.metadata_sheet = self.workbook.create_sheet("Metadata")
        self._write_metadata_sheet()

    def _write_metadata_sheet(self):
        self.metadata_sheet.delete_rows(1, self.metadata_sheet.max_row)
        for key, value in self.metadata.items():
            self.metadata_sheet.append([key, json.dumps(value, ensure_ascii=False)])

    def update_metadata(self, updates):
        self.metadata.update(updates)
        self._write_metadata_sheet()

    def write(self, record):
        row = row_for_file(record)
        # Save on every row so a long experiment still leaves a usable file if
        # the app or device connection fails before close() is called.
        self.sheet.append([row.get(field, "") for field in FIELDNAMES])
        self.workbook.save(self.filename)

    def close(self):
        self.workbook.save(self.filename)
